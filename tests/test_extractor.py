"""Unit tests for DocumentTextExtractor."""

from pathlib import Path
import docx
import pypdf
import pytest

from ordinale.extractor import DocumentTextExtractor, ExtractedDocument


@pytest.fixture
def extractor() -> DocumentTextExtractor:
    return DocumentTextExtractor(max_chars=300)


def test_is_supported(extractor: DocumentTextExtractor):
    assert extractor.is_supported("contract.docx")
    assert extractor.is_supported("statement.pdf")
    assert extractor.is_supported("article.html")
    assert extractor.is_supported("notes.md")
    assert extractor.is_supported("readme.txt")
    assert not extractor.is_supported("installer.exe")
    assert not extractor.is_supported("archive.zip")


def test_extract_missing_file(extractor: DocumentTextExtractor, tmp_path: Path):
    missing_path = tmp_path / "non_existent.docx"
    doc = extractor.extract(missing_path)
    assert doc.file_type == "missing"
    assert "File does not exist" in (doc.extraction_error or "")


def test_extract_plain_text_and_markdown(extractor: DocumentTextExtractor, tmp_path: Path):
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Meeting Notes\nDiscussed Q3 budget and CRA tax filing deadlines.", encoding="utf-8")

    doc = extractor.extract(txt_file)
    assert doc.file_type == "text"
    assert "Meeting Notes" in doc.text_snippet
    assert "CRA tax filing" in doc.text_snippet
    assert "notes.txt" in doc.prompt_text


def test_extract_html(extractor: DocumentTextExtractor, tmp_path: Path):
    html_file = tmp_path / "article.html"
    html_content = """
    <!DOCTYPE html>
    <html>
    <head>
        <title>Guide to University Admissions</title>
        <meta name="description" content="A comprehensive overview of college application requirements.">
    </head>
    <body>
        <script>console.log("ignore me");</script>
        <h1>Undergraduate Admissions 2026</h1>
        <p>Submit your high school transcripts and recommendation letters by November 15.</p>
    </body>
    </html>
    """
    html_file.write_text(html_content, encoding="utf-8")

    doc = extractor.extract(html_file)
    assert doc.file_type == "html"
    assert doc.metadata.get("title") == "Guide to University Admissions"
    assert "Undergraduate Admissions 2026" in doc.text_snippet
    assert "ignore me" not in doc.text_snippet
    assert "Guide to University Admissions" in doc.prompt_text


def test_extract_docx(extractor: DocumentTextExtractor, tmp_path: Path):
    docx_file = tmp_path / "sample_assignment.docx"
    doc_builder = docx.Document()
    doc_builder.core_properties.title = "CS101 Lab 3 Assignment"
    doc_builder.core_properties.author = "Jane Doe"
    doc_builder.add_heading("Lab 3: Binary Search Trees", level=1)
    doc_builder.add_paragraph("Implement insert, delete, and in-order traversal in Python.")

    # Add table with some details
    table = doc_builder.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Deliverable"
    table.cell(0, 1).text = "Weight"
    table.cell(1, 0).text = "bst.py"
    table.cell(1, 1).text = "10%"

    doc_builder.save(str(docx_file))

    doc = extractor.extract(docx_file)
    assert doc.file_type == "docx"
    assert doc.metadata.get("title") == "CS101 Lab 3 Assignment"
    assert doc.metadata.get("author") == "Jane Doe"
    assert "Binary Search Trees" in doc.text_snippet
    assert "insert, delete" in doc.text_snippet
    assert "bst.py" in doc.text_snippet
    assert "Jane Doe" in doc.prompt_text


def test_extract_pdf(extractor: DocumentTextExtractor, tmp_path: Path):
    pdf_file = tmp_path / "cra_notice.pdf"

    # Generate minimal valid PDF with pypdf
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    # pypdf Writer can set metadata
    writer.add_metadata({
        "/Title": "Notice of Assessment 2025",
        "/Author": "Canada Revenue Agency",
    })
    with open(pdf_file, "wb") as f:
        writer.write(f)

    doc = extractor.extract(pdf_file)
    assert doc.file_type == "pdf"
    assert doc.metadata.get("pages") == 1
    assert "Notice of Assessment 2025" in (doc.metadata.get("title") or "")


def test_extract_binary_fallback(extractor: DocumentTextExtractor, tmp_path: Path):
    bin_file = tmp_path / "installer.exe"
    bin_file.write_bytes(b"\x4D\x5A\x90\x00\x03\x00")

    doc = extractor.extract(bin_file)
    assert doc.file_type == "binary"
    assert doc.file_size_bytes == 6
    assert doc.metadata.get("extension") == ".exe"


def test_extract_rtf(extractor: DocumentTextExtractor, tmp_path: Path):
    rtf_file = tmp_path / "agreement.rtf"
    rtf_content = r"""{\rtf1\ansi\ansicpg1252\deff0\deflang1033{\fonttbl{\f0\fnil\fcharset0 Arial;}}
{\*\generator Riched20 10.0.19041}\viewkind4\uc1
\pard\sa200\sl276\slmult1\b\f0\fs24 Residential Tenancy Agreement\b0\par
This agreement is entered into on January 1, 2026.\par
Monthly rent: \$1,800.\par
}"""
    rtf_file.write_text(rtf_content, encoding="utf-8")

    doc = extractor.extract(rtf_file)
    assert doc.file_type == "rtf"
    assert "Residential Tenancy Agreement" in doc.text_snippet
    assert "Monthly rent: $1,800." in doc.text_snippet
    assert r"\rtf1" not in doc.text_snippet
    assert r"\fonttbl" not in doc.text_snippet
    assert r"\generator" not in doc.text_snippet
    assert "agreement.rtf" in doc.prompt_text
